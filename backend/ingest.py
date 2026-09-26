import os
import uuid

from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook

from embedding import create_embeddings
from vector_store import client, COLLECTION_NAME


DOCUMENTS_FOLDER = os.path.abspath(
    os.path.join(
        os.path.dirname(__file__),
        "..",
        "data",
        "documents"
    )
)


# --------------------------------
# PDF
# --------------------------------

def load_pdf(file_path):

    texts = []

    reader = PdfReader(file_path)

    for page_number, page in enumerate(reader.pages):

        text = page.extract_text()

        if text and text.strip():

            texts.append({
                "text": text.strip(),
                "source": os.path.basename(file_path),
                "type": "pdf",
                "page": page_number + 1
            })

    return texts


# --------------------------------
# WORD
# --------------------------------

def load_word(file_path):

    document = Document(file_path)

    paragraphs = []

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:

            paragraphs.append(text)

    text = "\n".join(paragraphs)

    if not text:
        return []

    return [{
        "text": text,
        "source": os.path.basename(file_path),
        "type": "word"
    }]


# --------------------------------
# EXCEL
# --------------------------------

def load_excel(file_path):

    workbook = load_workbook(
        file_path,
        data_only=True
    )

    documents = []

    for sheet in workbook.worksheets:

        rows = list(
            sheet.iter_rows(
                values_only=True
            )
        )

        if not rows:
            continue

        headers = [
            str(value).strip()
            if value is not None
            else ""
            for value in rows[0]
        ]

        for row_number, row in enumerate(rows[1:], start=2):

            values = list(row)

            if not any(
                value is not None
                for value in values
            ):
                continue

            row_data = []

            for header, value in zip(headers, values):

                if header and value is not None:

                    row_data.append(
                        f"{header}: {value}"
                    )

            if not row_data:
                continue

            text = "\n".join(row_data)

            documents.append({
                "text": text,
                "source": os.path.basename(file_path),
                "type": "excel",
                "sheet": sheet.title,
                "row": row_number
            })

    return documents


# --------------------------------
# LOAD ALL DOCUMENTS
# --------------------------------

def load_documents():

    documents = []

    for filename in os.listdir(DOCUMENTS_FOLDER):

        file_path = os.path.join(
            DOCUMENTS_FOLDER,
            filename
        )

        if filename.lower().endswith(".pdf"):

            documents.extend(
                load_pdf(file_path)
            )

        elif filename.lower().endswith(".docx"):

            documents.extend(
                load_word(file_path)
            )

        elif filename.lower().endswith(".xlsx"):

            documents.extend(
                load_excel(file_path)
            )

    return documents


# --------------------------------
# STORE IN QDRANT
# --------------------------------

def ingest():

    documents = load_documents()

    print(
        f"Loaded {len(documents)} document sections."
    )

    if not documents:

        print("No documents found.")

        return

    texts = [
        document["text"]
        for document in documents
    ]

    print("Creating embeddings...")

    embeddings = create_embeddings(texts)

    print(
        f"Created {len(embeddings)} embeddings."
    )

    points = []

    for document, embedding in zip(
        documents,
        embeddings
    ):

        points.append({

            "id": str(uuid.uuid4()),

            "vector": embedding,

            "payload": {
                "text": document["text"],
                "source": document["source"],
                "type": document["type"],
                "sheet": document.get("sheet"),
                "row": document.get("row"),
                "page": document.get("page")
            }
        })

    client.upsert(
        collection_name=COLLECTION_NAME,
        points=points
    )

    print(
        f"Stored {len(points)} vectors in Qdrant."
    )


if __name__ == "__main__":

    ingest()